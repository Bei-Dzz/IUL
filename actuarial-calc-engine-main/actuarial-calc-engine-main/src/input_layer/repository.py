"""
Input repository implementations for Excel and PostgreSQL data sources.
"""

import pandas as pd
from sqlalchemy import create_engine, text
from typing import Any, Dict, List, Union
import openpyxl

from ..core.base import RepositoryReader
from ..core.exceptions import InputError
from ..core.logging import logger


class ExcelRepository(RepositoryReader):
    """Repository for reading data from Excel files."""

    def __init__(self, file_path: str):
        """
        Initialize Excel repository.

        Args:
            file_path: Path to the Excel file
        """
        self.file_path = file_path
        self._workbook = None
        self._sheet_mappings: Dict[str, str] = {}

        # Load workbook and discover sheets
        try:
            self._workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            # Create default mappings (sheet name -> sheet name)
            for sheet_name in self._workbook.sheetnames:
                self._sheet_mappings[sheet_name.lower()] = sheet_name
        except Exception as e:
            raise InputError(f"Failed to load Excel file {file_path}: {e}")

    def read_data(self, key: str) -> Union[pd.DataFrame, Dict[str, Any], Any]:
        """
        Read data from Excel sheet by key.

        Args:
            key: Sheet name (case-insensitive)

        Returns:
            DataFrame containing the sheet data
        """
        if not self._workbook:
            raise InputError("Excel workbook not loaded")

        sheet_key = key.lower()
        if sheet_key not in self._sheet_mappings:
            available_keys = list(self._sheet_mappings.keys())
            raise InputError(f"Sheet '{key}' not found. Available sheets: {available_keys}")

        sheet_name = self._sheet_mappings[sheet_key]

        try:
            # Read sheet into DataFrame
            df = pd.read_excel(self.file_path, sheet_name=sheet_name, engine='openpyxl')
            logger.debug(f"Read {len(df)} rows from Excel sheet '{sheet_name}'")
            return df
        except Exception as e:
            raise InputError(f"Failed to read sheet '{sheet_name}': {e}")

    def list_available_keys(self) -> List[str]:
        """List all available sheet names."""
        return list(self._sheet_mappings.keys())


class PostgreSQLRepository(RepositoryReader):
    """Repository for reading data from PostgreSQL database."""

    def __init__(self, connection_string: str, table_mappings: Dict[str, str] = None):
        """
        Initialize PostgreSQL repository.

        Args:
            connection_string: Database connection string
            table_mappings: Optional mapping of keys to table names
        """
        self.connection_string = connection_string
        self.table_mappings = table_mappings or {}
        self._engine = None

        try:
            self._engine = create_engine(connection_string)
            # Test connection
            with self._engine.connect() as conn:
                conn.execute(text("SELECT 1"))
        except Exception as e:
            raise InputError(f"Failed to connect to PostgreSQL database: {e}")

    def read_data(self, key: str) -> Union[pd.DataFrame, Dict[str, Any], Any]:
        """
        Read data from database table by key.

        Args:
            key: Table key (from mappings or direct table name)

        Returns:
            DataFrame containing the table data
        """
        if not self._engine:
            raise InputError("Database connection not established")

        table_name = self.table_mappings.get(key, key)

        try:
            query = f"SELECT * FROM {table_name}"
            df = pd.read_sql(query, self._engine)
            logger.debug(f"Read {len(df)} rows from table '{table_name}'")
            return df
        except Exception as e:
            raise InputError(f"Failed to read table '{table_name}': {e}")

    def list_available_keys(self) -> List[str]:
        """List all available table keys."""
        if not self._engine:
            return []

        try:
            # Get all table names from information_schema
            query = """
            SELECT table_name
            FROM information_schema.tables
            WHERE table_schema = 'public'
            """
            with self._engine.connect() as conn:
                result = conn.execute(text(query))
                table_names = [row[0] for row in result.fetchall()]

            # Include both mapped keys and direct table names
            available_keys = list(self.table_mappings.keys()) + table_names
            return list(set(available_keys))  # Remove duplicates
        except Exception as e:
            logger.warning(f"Failed to list available tables: {e}")
            return list(self.table_mappings.keys())